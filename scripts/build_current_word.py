"""Editable A4 handbook from current inputs; no independent research or stale prose."""
import json, os, re
from collections import Counter
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

base=Path(os.environ['FRIESGO_OUTPUT_DIR'])
accounts=json.loads((base/'canonical_accounts.json').read_text())
model=json.loads((base/'operating_model.json').read_text())
doc=Document();s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.7);s.bottom_margin=Cm(1.6);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8)
s.header_distance=Cm(.7);s.footer_distance=Cm(.7)
for name,size in [('Normal',10.5),('Title',22),('Subtitle',11),('Heading 1',16),('Heading 2',12.5),('Heading 3',11)]:
    style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
    style.paragraph_format.space_after=Pt(5)
    style.paragraph_format.line_spacing=Pt(size*1.4)
    if name.startswith('Heading'):
        style.paragraph_format.keep_with_next=True;style.paragraph_format.space_before=Pt(9)
for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)

def para(text='',style=None):
    p=doc.add_paragraph(style=style)
    for part in re.split(r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*)',text):
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
        if m:
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=m[1];r.append(t);h.append(r);p._p.append(h)
        elif part:
            r=p.add_run(part[2:-2] if part.startswith('**') else part);r.bold=part.startswith('**')
    return p

def heading(text,level=1,page=False):
    clean=re.sub(r'[：:，、“”/（）()]',' ',text)
    p=para(clean,'Heading '+str(level))
    if page:p.paragraph_format.page_break_before=True
    return p

def markdown(file,skip_first=True):
    lines=file.read_text().splitlines()
    for i,line in enumerate(lines):
        if not line.strip() or skip_first and i==0:continue
        if line.startswith('#'):
            n=len(line)-len(line.lstrip('#'));heading(line[n:].strip(),min(n,3))
        else:para(line)

def letter(subject,body):
    para('**Subject：**'+subject)
    for text in body.split('\n\n'):
        p=para(text);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=Pt(14)

head=s.header.paragraphs[0];head.text='FriesGo Global  海外渠道开发';head.alignment=WD_ALIGN_PARAGRAPH.RIGHT
head.style=doc.styles['Normal'];head.runs[0].font.size=Pt(8)
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('FriesGo  v'+model['version']+'  |  ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
para('海外阀门渠道开发审核与执行手册','Title')
para('FriesGo Global  20家候选账户及逐家开发信','Subtitle')
para('报告日期 '+model['report_date']+'  原来源核验日期逐家保留')
counts=Counter(a['review']['queue'] for a in accounts)
para(f"本手册重审市场选择、经销商进入条件、产品可审核性和销售推进。现有{len(accounts)}家均为候选账户，建议{counts.get('P1 小批验证',0)}家小批验证、{counts.get('P2 条件验证',0)}家先补条件、{counts.get('P3 暂缓',0)}家暂缓；尚未确认采购职权、当前需求或正式报价条件。全部开发信为草稿，未发送。")
para('**使用顺序：**先读审核结论与推进手册，补一款产品的关键资料，再按当前P队列验证。20家档案保留公开联系、原事实和项目边界，随后给出首选语言及中文稿；越南账户另保留英文母稿。')
para('**与旧版关系：**原市场研究和已验收快照完整保留。旧A/B与分数只供研究追溯，当前行动以本版进入和停止条件为准。打印表仍是一张A4横版，首列只有公司名，160个操作格留空。')
heading('第一部分 项目审核与优化',page=True)
markdown(base/'01_项目审核与优化.md')
heading('第二部分 业务员推进与复盘',page=True)
markdown(base/'02_业务员推进手册.md')
heading('产品资料逐项清单',2)
markdown(base/'08_产品资料卡.md')
for number,item,status,owner,when in model['technical_fields']:
    para('**'+number+' '+item+'：**'+status+'。责任来源：'+owner+'；使用时点：'+when+'。')
heading('正式报价条件',2)
para('逐项确认规格、数量与时间、买方职权、目的地与收货主体、Incoterm、付款、适用准入与文件、包装运费、授权成本、批准及有效期。只有适用性评估可有依据地说明无特定文件要求；规格、成本、职权和交易条件不能用N/A免除。当前所有账户尚未齐备。')
heading('第三部分 当前20家队列索引',page=True)
para('越南、马来西亚和阿联酋是本案例的验证范围，不是全球市场最优排名。详细八国贸易和渠道来源见随包03原市场研究。HS848180只作宽口径背景；2026月度缺失不作零值或同比推断。')
for label in ['P1 小批验证','P2 条件验证','P3 暂缓']:
    heading(label,2)
    for a in accounts:
        if a['review']['queue']==label:para(a['id']+' '+a['display_name']+'  '+a['country']+'；'+a['review']['reason'])

for a in accounts:
    r=a['review'];heading(a['id']+' '+a['display_name'],page=True)
    para(a['name']);para(a['country']+'；'+a['role']+'；'+r['queue'])
    heading('公开联系与证据',2)
    for label,key in [('官网','website'),('业务邮箱','email'),('电话及已核即时渠道','phone'),('地址','address'),('联系人及转介岗位','person'),('备用渠道','alt')]:
        para('**'+label+'：**'+a[key])
    for label,value in [('公开事实',a['fact']),('历史项目或应用',a['scenario']),('项目边界',a['project_status']),('产品位置',r['entry_product']),('合作障碍',a['risk'])]:para('**'+label+'：**'+value)
    heading('本次判断和推进条件',2)
    for label,key in [('研究判断','reason'),('首封目标','next_commitment'),('进入下一步','unlock'),('停止条件','stop')]:para('**'+label+'：**'+r[key])
    para('**阶段：**'+r['stage']+'；采购职权与当前需求未确认；'+r['outreach_status'])
    para('**来源：**[产品与业务]('+a['product_url']+')；[联系与身份]('+a['contact_url']+')'+('；[项目补充]('+a['project_url']+')' if a['project_url'] else ''))
    para('原来源核验 '+a['checked_on']+'；判断复审 '+r['reviewed_on']+'。公开联系方式不代表投递成功。')
    heading(a['id']+' 开发信草稿',page=True)
    para('公开收件入口 '+a['email']+'。'+r['queue']+'；发送前核入口、签名及本页问题。')
    heading(a['language']+'首封',2);letter(a['local_subject'],a['local'])
    heading('中文译文',2)
    for tx in a['chinese'].split('\n\n'):para(tx)
    para('**下一步：**'+a['next_action'])

vietnamese=[a for a in accounts if a['language']=='越南语']
if vietnamese:
    heading('附录 越南账户英文母稿',page=True)
    for index,a in enumerate(vietnamese):
        heading(a['id']+' '+a['display_name'],2,page=index>0 and index%2==0);letter(a['subject'],a['english'])
doc.core_properties.title='FriesGo 海外阀门渠道开发审核与执行手册'
doc.core_properties.author='FriesGo Global'
out=base/'FriesGo_审核优化与20家客户开发手册_v2.docx';doc.save(out)
print(out.name)
