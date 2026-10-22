import json, re, os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(os.environ['FRIESGO_OUTPUT_DIR']).resolve()
OUT = BASE / 'FriesGo_海外渠道市场调研与20家客户开发方案_2026-10-23.docx'
accounts = json.loads((BASE/'canonical_accounts.json').read_text())
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.72), Inches(.72)
sec.header_distance, sec.footer_distance = Inches(.25), Inches(.25)

def font(run, size=None, bold=None):
    run.font.name = 'Arial'
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for k in ('ascii','hAnsi','cs'): rf.set(qn('w:'+k),'Arial')
    rf.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    run.font.color.rgb = RGBColor(0,0,0)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold

for nm, sz in [('Normal',11),('Title',24),('Subtitle',12),('Heading 1',17),('Heading 2',14),('Heading 3',12),('Heading 4',11)]:
    s=doc.styles[nm]; s.font.name='Arial'; s.font.size=Pt(sz); s.font.color.rgb=RGBColor(0,0,0)
    rf=s.element.get_or_add_rPr().get_or_add_rFonts(); rf.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    pf=s.paragraph_format; pf.space_after=Pt(6); pf.line_spacing=Pt(sz*1.42)
    if nm.startswith('Heading'): pf.space_before=Pt(12); pf.keep_with_next=True
doc.styles['Normal'].paragraph_format.widow_control=True
for e in doc.styles.element.xpath('.//w:pBdr'):
    e.getparent().remove(e)

def link(p, label, url):
    if not url: return
    if url.startswith('/'): url=Path(url).as_uri()
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    co=OxmlElement('w:color'); co.set(qn('w:val'),'245A81'); rp.append(co)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rp.append(u); r.append(rp)
    t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)

def inline(p, text):
    pat=r'(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for s in re.split(pat,text):
        if not s: continue
        m=re.match(r'\[([^\]]+)\]\(([^\)]+)\)',s)
        if m: link(p,m[1],m[2])
        else:
            bold=s.startswith('**'); t=s[2:-2] if bold else s.strip('`')
            font(p.add_run(t),bold=bold)

def para(text='', style=None, small=False):
    p=doc.add_paragraph(style=style); inline(p,text)
    if small:
        for r in p.runs: font(r,10)
    return p

page_pending=False
def new_page():
    global page_pending
    page_pending=True

def heading(text,level=1):
    global page_pending
    # Keep title and headings descriptive and free of ornamental punctuation.
    text=re.sub(r'[：:，“”「」/]', ' ', text).strip()
    p=doc.add_paragraph(text, 'Heading '+str(level))
    if page_pending:
        p.paragraph_format.page_break_before=True
        page_pending=False
    for r in p.runs: font(r,bold=True)
    return p

def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    for j,h in enumerate(headers): t.rows[0].cells[j].text=h
    trPr=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); trPr.append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for j,x in enumerate(row): cells[j].text=str(x)
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders); mar=OxmlElement('w:tcMar')
            for edge,val in [('top',70),('bottom',70),('left',90),('right',90)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF0F4');pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(14)
                for r in p.runs: font(r,10.5,bold=i==0)
            tr=row._tr.get_or_add_trPr(); ns=OxmlElement('w:cantSplit');tr.append(ns)
    para('').paragraph_format.space_after=Pt(0)
    return t

header=sec.header.paragraphs[0];header.text='FriesGo Global  海外阀门渠道开发';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs: font(r,9)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(f.add_run('第 '),9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
font(f.add_run(' 页'),9)

para('海外阀门渠道市场调研与客户开发', 'Title')
para('FriesGo Global  渠道账户及独立开发信', 'Subtitle')
para('核验基准日 2026年10月23日',small=True)
para('本文件先比较市场，再识别海外分销网络，最后逐家核实商户和联系入口。首批选择越南7家、马来西亚6家、阿联酋7家，共20家，重点开发能够向工程和楼宇客户转售水系统阀门的进口商、分销商及配套供应商。')
para('产品范围为手动平衡阀、自力式压差控制阀、自力式流量控制阀。客户档案包含公开业务邮箱、电话、地址、渠道角色、项目或应用依据、合作障碍及逐家开发信。7家越南客户附越南语正文，其余采用英文，所有客户均附中文译文。')
para('**阅读提示：**这些是已核公司与渠道相关性的潜在客户，尚未确认采购权限、当前询价或购买意愿。历史项目只用于说明渠道经历，未把历史供货写成正在采购。开发信均为草稿，尚未发送。')
heading('文件结构',2)
para('第一部分说明产品边界、八国比较、六个渠道网络、入选依据和跟进方法。第二部分提供20家客户索引。第三部分逐家给出联系资料、证据、开发判断和可复制的信件。来源名称可点击访问；本地产品母表的链接需在原工作电脑打开。')

new_page()
heading('第一部分 市场调研与渠道选择依据')
md=(BASE/'01_市场调研与渠道选择依据.md').read_text().split('## 交付索引')[0]
lines=md.splitlines(); i=1
while i<len(lines):
    ln=lines[i].strip()
    if not ln: i+=1;continue
    if ln.startswith('|'):
        chunk=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in row): chunk.append(row)
            i+=1
        h,rs=chunk[0],chunk[1:]
        if h[0]=='市场': table(h,rs,[1.08,.65,.65,.65,.8,2.97])
        else:
            for row in rs:
                heading(row[0]+' '+row[1] if h[0]=='渠道ID' else row[0],3)
                for k,v in zip(h[2:] if h[0]=='渠道ID' else h[1:],row[2:] if h[0]=='渠道ID' else row[1:]):
                    para('**'+k+'：**'+v)
        continue
    if ln.startswith('#'):
        n=len(ln)-len(ln.lstrip('#'));heading(ln[n:].strip(),min(n,4))
    elif ln.startswith('- '): para('• '+ln[2:])
    else: para(ln)
    i+=1

new_page();heading('第二部分 20家客户快速索引')
para('A组优先确认选品信息，B组先核合作障碍。分数为公开资料研究判断，并非成交概率。所有账户的采购权限与当前采购需求均未确认。')
short=[a.get('display_name',a['name']) for a in accounts]
table(['编号','市场','公司简称','组别 分数','重点产品'],[[a['id'],a['country'],s,a['priority']+' '+str(a['score']),a['product']] for a,s in zip(accounts,short)],[.48,.8,1.6,.78,3.14])
para('评分分项顺序为：产品及应用25、渠道角色20、市场匹配15、商业品类信号15、采购岗位可见度10、可联系性10、证据质量及新鲜度5。80分及以上为A组，65至79分为B组。品类在售不等于当前采购。')

for a in accounts:
    new_page()
    heading(a['id']+' '+a['name'])
    para(a['country']+'  '+a['priority']+'组  '+str(a['score'])+'/100  '+a['role']+'  渠道 '+a['channel'],small=True)
    heading('联系资料',2)
    t=table(['字段','公开资料'],[
        ('首选业务邮箱',a['email']),('电话及即时渠道',a['phone']),('地址',a['address']),
        ('联系人及转介岗位',a['person']),('备用渠道',a['alt']),('官网',a['website'])],[1.12,5.68])
    para('采购权限未确认。'+a['verification']+'。',small=True)
    heading('入选依据与首封目标',2)
    for k,v in [('已核事实',a['fact']),('匹配产品',a['product']),('项目或应用',a['scenario']),('项目证据状态',a['project_status']),('开发判断',a['angle']),('合作障碍',a['risk']),('首封目标',a['ask'])]:
        p=para('**'+k+'：**'+v);p.paragraph_format.space_after=Pt(4)
    p=para('**来源：**')
    link(p,'产品及业务依据',a['product_url']);p.add_run('  ');link(p,'联系及身份依据',a['contact_url'])
    if a['project_url']:
        p.add_run('  ');link(p,'项目或补充依据',a['project_url'])
    para('核验日 '+a['checked_on']+'  分项评分 '+ ' / '.join(map(str,a['components']))+'。公司自述不等于独立业绩审计。',small=True)

    new_page()
    heading(a['id']+' 开发信与联系动作')
    para('收件入口：'+a['email']+'。首次不附通用公司宣传附件；请按实际发件账户补入签名。',small=True)
    def letter(label,subject,body):
        heading(label,2)
        p=para('**Subject：**'+subject);p.paragraph_format.space_after=Pt(4)
        for tx in body.split('\n\n'):
            p=para(tx);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=Pt(14.5)
            # Source letters retain all original wording and paragraph boundaries.
    if a['language']=='越南语':
        letter('越南语首封正文',a['local_subject'],a['local'])
        para('英文母稿见文末附录中对应客户编号。',small=True)
    else: letter('英文首封正文',a['subject'],a['english'])
    heading('中文译文',2)
    for tx in a['chinese'].split('\n\n'):
        p=para(tx);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=Pt(14.5)
    heading('下一步联系动作',2)
    para(a['next_action'],small=True)

new_page();heading('附录 越南客户英文母稿')
para('以下保留7家越南客户的英文母稿，可供内部审核或在客户要求英语沟通时使用。默认首封正文及中文译文仍在对应客户档案中。')
appendix_count=0
for a in accounts:
    if a['language']!='越南语': continue
    if appendix_count and appendix_count%2==0: new_page()
    heading(a['id']+' '+short[accounts.index(a)],2)
    para('**Subject：**'+a['subject'])
    for tx in a['english'].split('\n\n'):
        p=para(tx);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=Pt(14.5)
    appendix_count+=1

doc.core_properties.title='FriesGo 海外阀门渠道市场调研与20家客户开发方案'
doc.core_properties.subject='三类水系统阀门的市场筛选 分销渠道 客户联系方式及独立开发信'
doc.core_properties.author='FriesGo Global'
doc.save(OUT)
print(OUT)
print(json.dumps({'accounts':len(accounts),'letters_english':len(accounts),'letters_vietnamese':sum(a['language']=='越南语' for a in accounts),'letters_chinese':len(accounts),'tables':len(doc.tables)},ensure_ascii=False))
